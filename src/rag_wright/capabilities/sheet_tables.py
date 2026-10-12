"""PS-26 (ADR-0136): glue a spreadsheet table back together across its blank rows.

docling cuts a sheet's table at every blank row, and every piece's first row is then read as a header, so the first
record after each blank row was lost. After the parse, a piece directly below a table, in the same columns, with
only blank rows between, is decided by what its first row is:

  data   : glued on (its rows appended to the table above, its sheet rectangle kept as another `prov` entry so a
           sheet cell still maps to its table row). Decided by the number or date columns (every data row above is a
           number or date there, the header is not, and the first row is too) or, when there are none, by formatting
           (the header is formatted differently from the data rows, and the first row is formatted like the data).
  header : left as its own table -- the same header repeated, a header formatted like the one above, text where
           the data above has numbers or dates, or a table above whose top row does not name every column its rows
           use (a form's section label, not a column header), or a first row that does not fill the columns the
           rows above fill (a row of that table would).
  unsure : neither signal decides (a text-only table with no formatting): left as its own table, and reported so
           `evaluate_ingestion` can warn.

A piece that is not directly below a table in the same columns (a title, a side table, a form's label column) is
never touched.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from typing import Any, Literal, Optional

_TYPED = re.compile(r"^[-+]?(\d[\d,]*\.?\d*|\.\d+)\s*%?$|^\d{1,4}[-/.]\d{1,2}[-/.]\d{1,4}")


@dataclass(frozen=True)
class Continuation:
    """A table piece directly below another in the same columns: `verdict` says whether its first row is data."""

    sheet: str
    upper: str  # the table above (self-ref)
    lower: str  # the piece below (self-ref)
    sheet_row: int  # 1-based sheet row of the piece's first row
    verdict: Literal["data", "header", "unsure"]


def _text(cell: Any) -> str:
    return " ".join((cell.text or "").split())


def _typed(value: str) -> bool:
    return bool(value) and bool(_TYPED.match(value))


def _box(table: Any) -> Any:
    return table.prov[0].bbox


def _bottom(table: Any) -> float:
    """The table's last sheet row (exclusive): a glued table ends where its last block ends."""
    return max(p.bbox.b for p in table.prov)


def _sheet_rows(table: Any) -> list[int]:
    """The 1-based sheet row of each grid row, through each block's rectangle (`prov`), in grid order."""
    return [r + 1 for p in table.prov for r in range(int(p.bbox.t), int(p.bbox.b))]


def _sheet_tables(document: Any) -> dict[str, list]:
    from rag_wright.ingestion.tables import _sheet_of, visible_tables

    by_sheet: dict[str, list] = {}
    for t in visible_tables(document):
        sheet = _sheet_of(t, document)
        if sheet is not None and t.prov and t.data.grid:
            by_sheet.setdefault(sheet, []).append(t)
    return by_sheet


def _style(ws: Any, sheet_row: int, first_col: int, last_col: int) -> Optional[tuple[bool, bool]]:
    """(bold, filled) by majority over the row's non-empty cells; None when the row has none."""
    cells = [ws.cell(row=sheet_row, column=c) for c in range(first_col, last_col + 1)]
    cells = [c for c in cells if c.value not in (None, "")]
    if not cells:
        return None
    bold = sum(bool(c.font and c.font.b) for c in cells) * 2 > len(cells)
    filled = sum(bool(c.fill and c.fill.fill_type) for c in cells) * 2 > len(cells)
    return bold, filled


def _verdict(upper: Any, lower: Any, ws: Any) -> Literal["data", "header", "unsure"]:
    grid = [[_text(c) for c in row] for row in upper.data.grid]
    header, data = grid[0], grid[1:]
    first = [_text(c) for c in lower.data.grid[0]]
    if not data or first == header:
        return "header"
    used = {j for row in data for j, v in enumerate(row) if v}
    if any(not header[j] for j in used):
        return "header"  # the top row does not name the columns its rows use: a form section, not a database
    always = set.intersection(*({j for j, v in enumerate(row) if v} for row in data))
    filled = {j for j, v in enumerate(first) if v}
    if filled - used or always - filled:
        return "header"  # it does not fill the columns its would-be rows fill: not one of them
    typed = [j for j in range(len(header)) if all(_typed(r[j]) for r in data) and not _typed(header[j])]
    if typed:
        return "data" if all(_typed(first[j]) for j in typed) else "header"
    if ws is None:
        return "unsure"
    rows, left, right = _sheet_rows(upper), int(_box(upper).l) + 1, int(_box(upper).r)
    header_style = _style(ws, rows[0], left, right)
    data_styles = [s for s in (_style(ws, r, left, right) for r in rows[1:]) if s is not None]
    data_style = max(set(data_styles), key=data_styles.count) if data_styles else None
    first_style = _style(ws, int(_box(lower).t) + 1, left, right)
    if header_style is None or data_style is None or header_style == data_style:
        return "unsure"
    return "data" if first_style == data_style else "header" if first_style == header_style else "unsure"


def _workbook(data: Optional[bytes]) -> Any:
    if not data:
        return None
    try:
        import openpyxl

        return openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
    except Exception:  # noqa: BLE001 - not an OOXML workbook (CSV, ODS, legacy XLS): no formatting signal
        return None


def continuations(document: Any, data: Optional[bytes]) -> list[Continuation]:
    """Every table piece directly below another in the same columns, with only blank rows between, top to bottom
    per sheet, each with its verdict. `data` is the workbook's bytes (for formatting); None = numbers and dates
    only."""
    from rag_wright.corpus.document_parser import is_spreadsheet

    if not is_spreadsheet(document):
        return []
    book = _workbook(data)
    out: list[Continuation] = []
    for sheet, tables in _sheet_tables(document).items():
        ws = book[sheet] if book is not None and sheet in book.sheetnames else None
        for i, lower in enumerate(tables):
            lb = _box(lower)
            above = [t for t in tables[:i] if _box(t).l == lb.l and _box(t).r == lb.r and _bottom(t) <= lb.t]
            if not above:
                continue
            upper = max(above, key=_bottom)
            between = [t for t in tables if t is not upper and t is not lower and _box(t).t >= _bottom(upper)
                       and _bottom(t) <= lb.t and _box(t).l < lb.r and _box(t).r > lb.l]
            if between or len(upper.data.grid[0]) != len(lower.data.grid[0]):
                continue
            out.append(Continuation(sheet, upper.self_ref, lower.self_ref, int(lb.t) + 1,
                                    _verdict(upper, lower, ws)))
    return out


def merge_continuations(document: Any, data: Optional[bytes]) -> bool:
    """Glue every `data` continuation onto the table above it (repeatedly, so a table cut at several blank rows
    comes back whole). Returns whether anything changed."""
    changed = False
    while True:
        step = next((c for c in continuations(document, data) if c.verdict == "data"), None)
        if step is None:
            return changed
        upper = _resolve(document, step.upper)
        lower = _resolve(document, step.lower)
        offset = upper.data.num_rows
        for cell in lower.data.table_cells:
            upper.data.table_cells.append(cell.model_copy(update={
                "start_row_offset_idx": cell.start_row_offset_idx + offset,
                "end_row_offset_idx": cell.end_row_offset_idx + offset, "column_header": False}))
        upper.data.num_rows += lower.data.num_rows
        upper.prov.extend(lower.prov)  # each glued block's sheet rectangle, in grid order
        document.delete_items(node_items=[lower])
        changed = True


def _resolve(document: Any, self_ref: str) -> Any:
    return next(t for t in document.tables if t.self_ref == self_ref)
