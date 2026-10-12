"""EV-2: gold sets imported from an expert's own files (CSV, XLSX, JSONL) through a column mapping, exactly as
written: a title above the header, blank rows, stray spaces and odd header spelling do not lose or shift a row."""
from __future__ import annotations

import asyncio
import datetime as dt
import json
from pathlib import Path

import openpyxl
import pytest

from rag_wright.api import GoldImportError, GoldMapping, IngestSource, aload_gold_set, load_gold_set

MAP = GoldMapping(id="ID", input=["Premise", "Hypothesis"], expected="Label", group="Contract",
                  metadata=["Notes"])


def _xlsx(path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gold"
    for row in [["NLI gold set - reviewed by legal (v3)"], [],
                [" id ", " Premise \n", "hypothesis", "LABEL ", "Contract", "Notes"],
                ["a1", "  The term is\n2 years. ", "Term is two years.", "entailment ", "c-1", None],
                [],
                ["a2", "No fee applies.", "A fee applies.", "contradiction", "c-1", "checked twice"],
                ["a3", 12.5, "x", "neutral", "c-2", dt.date(2026, 10, 1)],
                [None, None, None, None, None, None],
                ["a4", "Either party may end it.", "Termination is mutual.", "", "c-3", None]]:
        ws.append(row)
    other = wb.create_sheet("Summary")
    other.append(["Total", 4])
    hidden = wb.create_sheet("Old")
    hidden.append(["ID", "Premise", "Hypothesis", "Label", "Contract", "Notes"])
    hidden.append(["z9", "old", "old", "neutral", "c-9", None])
    hidden.sheet_state = "hidden"
    wb.save(path)
    return path


def _map(**kw) -> GoldMapping:
    return MAP.model_copy(update=kw)


def test_a_messy_workbook_imports_every_row_exactly(tmp_path):
    gold = load_gold_set(_xlsx(tmp_path / "gold.xlsx"), MAP, name="nli", version="v3")
    assert [i.id for i in gold.items] == ["a1", "a2", "a3", "a4"]
    a1, a2, a3, a4 = gold.items
    assert a1.input == {"Premise": "The term is\n2 years.", "Hypothesis": "Term is two years."}
    assert a1.expected == "entailment" and a1.group == "c-1"
    assert a3.input["Premise"] == "12.5" and a3.metadata["Notes"] == "2026-10-01"
    assert a4.expected == ""  # an empty label is kept: "" means none
    assert a2.metadata == {"Notes": "checked twice", "source": "gold.xlsx: Gold row 6"}
    assert gold.local_only is True and gold.name == "nli" and gold.version == "v3"


def test_hidden_sheets_are_read_only_when_named(tmp_path):
    book = _xlsx(tmp_path / "gold.xlsx")
    assert "z9" not in [i.id for i in load_gold_set(book, MAP, name="n", version="1").items]
    assert [i.id for i in load_gold_set(book, MAP, name="n", version="1", sheet="Old").items] == ["z9"]


def test_a_missing_column_names_what_each_sheet_has(tmp_path):
    with pytest.raises(GoldImportError, match=r"missing Verdict -- .*Gold: id, Premise, hypothesis, LABEL"):
        load_gold_set(_xlsx(tmp_path / "gold.xlsx"), _map(expected="Verdict"), name="n", version="1")


def test_ids_default_to_where_the_row_is(tmp_path):
    gold = load_gold_set(_xlsx(tmp_path / "gold.xlsx"), _map(id=None), name="n", version="1")
    assert [i.id for i in gold.items] == ["Gold!4", "Gold!6", "Gold!7", "Gold!9"]


def test_a_csv_with_a_title_line_quoted_newlines_and_a_bom(tmp_path):
    path = tmp_path / "gold.csv"
    path.write_bytes(("﻿NLI gold\n\nID,Premise,Hypothesis,Label,Contract,Notes\n"
                      'a1,"The term is\n2 years.",Term is two years.,entailment,c-1,\n'
                      ",,,,,\n"
                      "a2, No fee. ,A fee applies., contradiction ,c-1,ok\n").encode("utf-8"))
    gold = load_gold_set(path, MAP, name="n", version="1")
    assert [(i.id, i.expected) for i in gold.items] == [("a1", "entailment"), ("a2", "contradiction")]
    assert gold.items[0].input["Premise"] == "The term is\n2 years." and gold.items[1].input["Premise"] == "No fee."
    assert gold.items[1].metadata["source"] == "gold.csv: row 6"


def test_jsonl_keeps_its_types(tmp_path):
    path = tmp_path / "gold.jsonl"
    rows = [{"ID": "a1", "Premise": "x", "Hypothesis": "y", "Label": ["entailment", "neutral"], "Contract": "c-1",
             "Notes": {"by": "legal"}},
            {"ID": "a2", "Premise": "x", "Hypothesis": "z", "Label": "neutral", "Contract": "c-2", "Notes": None}]
    path.write_text("\n".join(json.dumps(r) for r in rows) + "\n\n", encoding="utf-8")
    gold = load_gold_set(path, MAP, name="n", version="1")
    assert gold.items[0].expected == ["entailment", "neutral"] and gold.items[0].metadata["Notes"] == {"by": "legal"}
    assert gold.items[1].metadata["source"] == "gold.jsonl: line 2"


def test_several_expected_columns_become_one_record(tmp_path):
    path = tmp_path / "ext.csv"
    path.write_text("ID,Text,Party,Amount\nr1,Pay 500 to Acme,Acme,500\n", encoding="utf-8")
    gold = load_gold_set(path, GoldMapping(id="ID", input="Text", expected=["Party", "Amount"]), name="n",
                         version="1")
    assert gold.items[0].input == "Pay 500 to Acme" and gold.items[0].expected == {"Party": "Acme", "Amount": "500"}


def test_duplicate_ids_are_refused_by_name(tmp_path):
    path = tmp_path / "dup.csv"
    path.write_text("ID,Premise,Hypothesis,Label,Contract,Notes\na1,x,y,neutral,c,\na1,x,z,neutral,c,\n")
    with pytest.raises(GoldImportError, match="duplicate item ids: a1"):
        load_gold_set(path, MAP, name="n", version="1")


def test_a_row_without_its_id_is_refused_with_its_place(tmp_path):
    path = tmp_path / "noid.csv"
    path.write_text("ID,Premise,Hypothesis,Label,Contract,Notes\n,x,y,neutral,c,\n")
    with pytest.raises(GoldImportError, match="row 2.*no ID"):
        load_gold_set(path, MAP, name="n", version="1")


def test_provenance_can_come_from_a_column(tmp_path):
    path = tmp_path / "prov.csv"
    path.write_text("ID,Text,Label,Source\nr1,a,x, Silver \nr2,b,y,\n")
    gold = load_gold_set(path, GoldMapping(id="ID", input="Text", expected="Label", provenance="Source"),
                         name="n", version="1", provenance="hand")
    assert [i.provenance for i in gold.items] == ["silver", "hand"]
    path.write_text("ID,Text,Label,Source\nr1,a,x,guess\n")
    with pytest.raises(GoldImportError, match="row 2.*guess"):
        load_gold_set(path, GoldMapping(id="ID", input="Text", expected="Label", provenance="Source"),
                      name="n", version="1")


def test_an_uploaded_file_imports_from_its_bytes(tmp_path):
    data = _xlsx(tmp_path / "gold.xlsx").read_bytes()
    gold = load_gold_set(IngestSource(data=data, name="upload.xlsx"), MAP, name="n", version="1")
    assert len(gold.items) == 4 and gold.items[0].metadata["source"] == "upload.xlsx: Gold row 4"


def test_an_unknown_format_is_refused(tmp_path):
    (tmp_path / "gold.pdf").write_bytes(b"%PDF")
    with pytest.raises(GoldImportError, match="CSV, XLSX or JSONL"):
        load_gold_set(tmp_path / "gold.pdf", MAP, name="n", version="1")


def test_the_async_form_runs_inside_a_running_loop(tmp_path):
    book = _xlsx(tmp_path / "gold.xlsx")
    gold = asyncio.run(aload_gold_set(book, MAP, name="n", version="1"))
    assert [i.id for i in gold.items] == ["a1", "a2", "a3", "a4"]
