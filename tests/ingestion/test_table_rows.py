"""ING-7 (ADR-0124): the generic table-rows primitive -- every parsed table's rows as exact cells, read from the
parsed grid (whole even when the chunker split the table), with no domain knowledge."""
from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from rag_wright.api import IngestSource, TableRow, build_ingestion, parse_document, table_rows

from tests.corpus._ooxml_fixtures import make_xlsx


def _csv(tmp_path: Path, name: str, text: str):
    f = tmp_path / name
    f.write_text(text)
    return parse_document(f.stem, f, cache_dir=tmp_path / "cache")


def test_a_csv_table_comes_back_as_exact_rows(tmp_path):
    sd = _csv(tmp_path, "results.csv", 'Sr#,Sample,"Fabric, composition",GSM\n'
                                       '1,SDP 3101,"60% cotton, 40% poly",380\n2,SDP 3102,100% nylon,412\n')
    rows = table_rows(sd)
    assert [r.row_index for r in rows] == [1, 2]
    first = rows[0]
    assert first.columns == ["Sr#", "Sample", "Fabric, composition", "GSM"]
    assert first.values == ["1", "SDP 3101", "60% cotton, 40% poly", "380"]
    assert first.cell("Fabric, composition") == "60% cotton, 40% poly" and first.cell("missing") is None
    assert first.table_ref and first.sheet is None


def test_workbook_rows_carry_their_sheet_including_hidden_ones(tmp_path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    for r in (["Sr#", "Sample", "GSM"], [1, "SDP 1", 380], [2, "SDP 2", 412]):
        ws.append(r)
    hidden = wb.create_sheet("Trial log")
    for r in (["Sr.no", "Findings"], [1, "more stretch than the 75+40 trial"]):
        hidden.append(r)
    hidden.sheet_state = "hidden"
    path = tmp_path / "db.xlsx"
    wb.save(path)
    rows = table_rows(parse_document("db", path, cache_dir=tmp_path / "cache"))
    by_sheet = {r.sheet: r for r in rows}
    assert {r.sheet for r in rows} == {"Results", "Trial log"}
    assert by_sheet["Trial log"].cell("Findings") == "more stretch than the 75+40 trial"


def test_a_table_split_across_chunks_comes_back_whole(tmp_path):
    from rag_wright.capabilities.parsing import load_document
    from rag_wright.capabilities.rlm_chunking import StructuralBoundaryDiscoverer, chunk_texts

    lines = ["Sr#,Sample,Customer,Standard,Remarks"] + [
        f"{i},SDP {i},Northwind Apparel Ltd,EN 17092,trial note number {i} for the long-running wash test"
        for i in range(1, 3001)]
    sd = _csv(tmp_path, "big.csv", "\n".join(lines) + "\n")
    assert len(chunk_texts(load_document(sd.parsed), discoverer=StructuralBoundaryDiscoverer())) > 1  # really split
    rows = table_rows(sd)
    assert len(rows) == 3000 and all(len(r.values) == 5 for r in rows)
    assert rows[-1].values == ["3000", "SDP 3000", "Northwind Apparel Ltd", "EN 17092",
                               "trial note number 3000 for the long-running wash test"]


_FIXTURE = Path("tests/fixtures/table-bearing-contract.pdf")


@pytest.mark.skipif(not _FIXTURE.exists(), reason="PDF fixture not present")
def test_a_pdf_table_comes_back_with_its_page(tmp_path, monkeypatch):
    monkeypatch.setattr("rag_wright.capabilities.parsing._vlm_available", lambda: False)  # never a real VLM call
    rows = table_rows(parse_document("contract", _FIXTURE, cache_dir=tmp_path / "cache"))
    assert rows and all(r.page and r.page >= 1 for r in rows) and all(any(v for v in r.values) for r in rows)


def test_a_header_only_table_has_no_rows(tmp_path):
    assert table_rows(_csv(tmp_path, "empty.csv", "Sr#,Sample,GSM\n")) == []


def test_a_record_unit_carries_its_table_row_inside_the_builder(tmp_path):
    from tests.ingestion.test_builder import _FakeEmbedder, _ws

    rows = [["Sr#", "Sample", "Customer", "Standard", "Test", "Level"]] + [
        [i, f"SDP {3100 + i}", "Northwind", "EN 17092", "Tear", "AA"] for i in range(1, 6)]
    x = make_xlsx(tmp_path / "db.xlsx", rows, [])
    seen: list = []

    async def extractor(unit, *, source_doc_id):
        from rag_wright.api import KgNode, UnitExtraction

        seen.append(unit.table_row)
        return UnitExtraction(nodes=[KgNode("Record", "record_id", {
            "record_id": f"{source_doc_id}:{unit.index}", "span_id": unit.anchor.span_id,
            "confidence": "EXTRACTED"})])

    pipe = build_ingestion(extractor, embedder=_FakeEmbedder(), progress=lambda _l: None)
    asyncio.run(pipe.aingest(_ws(), [IngestSource(path=str(x), doc_id="db")], cache_dir=tmp_path / "c"))
    assert len(seen) == 5 and all(isinstance(r, TableRow) for r in seen)
    assert sorted(r.cell("Sample") for r in seen) == [f"SDP {3100 + i}" for i in range(1, 6)]
    assert all(r.row_index == int(r.cell("Sr#")) for r in seen)  # each unit got ITS row
