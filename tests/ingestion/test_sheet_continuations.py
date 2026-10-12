"""PS-26: a spreadsheet table with blank rows inside it is cut into pieces by the parse, and each piece's first row
was read as a header -- a record lost at every blank row. After the parse, a piece directly below a table in the
same columns is glued back on when its first row is DATA (number or date columns agree, or it is formatted like the
data rather than like the header); a header, a title, a different table or a form section stays apart. A text-only
table with no formatting cannot be decided from the file: it stays apart and `evaluate_ingestion` warns."""
from __future__ import annotations

from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from rag_wright.api import IngestSource, evaluate_ingestion, parse_document_bytes, table_rows
from tests.corpus._ooxml_fixtures import fake_pdf, make_xlsx, packager

H = ["ID", "Fabric", "GSM", "Result"]


def _book(path: Path, rows: list[list], bold_rows: tuple[int, ...] = ()) -> Path:
    """A one-sheet workbook; `bold_rows` are 0-based row numbers set in bold (a formatted header)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Results"
    for r in rows:
        ws.append(r)
    for i in bold_rows:
        for cell in ws[i + 1]:
            cell.font = Font(bold=True)
    wb.save(path)
    return path


def _rows(path: Path, cache: Path) -> list:
    return table_rows(parse_document_bytes("doc", path.name, path.read_bytes(), cache_dir=cache))


def _records(rows: list) -> list[tuple]:
    return [(tuple(r.columns), tuple(r.values)) for r in rows]


def test_number_columns_glue_a_database_back_across_blank_rows(tmp_path):
    book = _book(tmp_path / "db.xlsx", [H, ["1", "Jersey", 180, "pass"], ["2", "Rib", 220, "fail"], [],
                                        ["3", "Pique", 200, "pass"], ["4", "Fleece", 300, "pass"], [], [],
                                        ["5", "Terry", 350, "fail"]])
    rows = _rows(book, tmp_path / "c")
    assert [r.values[0] for r in rows] == ["1", "2", "3", "4", "5"]
    assert {tuple(r.columns) for r in rows} == {tuple(H)}
    assert len({r.table_ref for r in rows}) == 1


def test_formatting_glues_a_text_only_database(tmp_path):
    book = _book(tmp_path / "text.xlsx", [["Code", "Fabric", "Finish"], ["A", "Jersey", "soft"], [],
                                          ["B", "Rib", "matte"], ["C", "Pique", "gloss"]], bold_rows=(0,))
    rows = _rows(book, tmp_path / "c")
    assert [r.values[0] for r in rows] == ["A", "B", "C"]
    assert {tuple(r.columns) for r in rows} == {("Code", "Fabric", "Finish")}


def test_a_text_only_table_without_formatting_stays_apart_and_is_warned(tmp_path):
    book = _book(tmp_path / "plain.xlsx", [["Code", "Fabric", "Finish"], ["A", "Jersey", "soft"], [],
                                           ["B", "Rib", "matte"], ["C", "Pique", "gloss"]])
    rows = _rows(book, tmp_path / "c")
    assert len({r.table_ref for r in rows}) == 2  # today's split: undecidable from the file
    ev = evaluate_ingestion([str(book)], cache_dir=tmp_path / "e")
    assert any("plain.xlsx" in w and "may continue past a blank row" in w and "Results" in w for w in ev.warnings)
    assert ev.passed  # a warning, not a failure


def test_a_repeated_header_keeps_its_blocks_and_every_record(tmp_path):
    book = _book(tmp_path / "rep.xlsx", [H, ["1", "Jersey", 180, "pass"], [], H, ["2", "Rib", 220, "fail"]])
    rows = _rows(book, tmp_path / "c")
    assert _records(rows) == [(tuple(H), ("1", "Jersey", "180", "pass")), (tuple(H), ("2", "Rib", "220", "fail"))]


def test_a_title_above_a_table_stays_apart(tmp_path):
    book = _book(tmp_path / "title.xlsx", [["Lab results - Q3"], [], H, ["1", "Jersey", 180, "pass"],
                                           ["2", "Rib", 220, "fail"]])
    rows = _rows(book, tmp_path / "c")
    assert [r.values[0] for r in rows] == ["1", "2"] and {tuple(r.columns) for r in rows} == {tuple(H)}


def test_a_different_table_below_stays_apart(tmp_path):
    book = _book(tmp_path / "two.xlsx", [H, ["1", "Jersey", 180, "pass"], [],
                                         ["Supplier", "City", "Rating", "Since"], ["Acme", "Delhi", "A", 2019]])
    rows = _rows(book, tmp_path / "c")
    assert _records(rows) == [(tuple(H), ("1", "Jersey", "180", "pass")),
                              (("Supplier", "City", "Rating", "Since"), ("Acme", "Delhi", "A", "2019"))]


def test_a_formatted_header_below_starts_a_new_table(tmp_path):
    book = _book(tmp_path / "twotext.xlsx", [["Code", "Fabric"], ["A", "Jersey"], [], ["Supplier", "City"],
                                             ["Acme", "Delhi"]], bold_rows=(0, 3))
    rows = _rows(book, tmp_path / "c")
    assert _records(rows) == [(("Code", "Fabric"), ("A", "Jersey")), (("Supplier", "City"), ("Acme", "Delhi"))]


def test_form_sections_are_unchanged(tmp_path):
    form = [["Section A", "", ""], ["Buyer", "Acme", ""], ["Style", "T-100", ""], [],
            ["Section B", "", ""], ["Gauge", "28G", ""], ["Stitch", "2.8", ""]]
    book = _book(tmp_path / "form.xlsx", form)
    rows = _rows(book, tmp_path / "c")
    assert len({r.table_ref for r in rows}) == 2
    assert evaluate_ingestion([str(book)], cache_dir=tmp_path / "e").warnings == []  # a section label, no warning


def test_an_embedded_file_in_a_glued_row_maps_to_that_row(tmp_path):
    rows = [H, ["1", "Jersey", 180, "pass"], [], ["2", "Rib", 220, "fail"], ["3", "Pique", 200, "pass"]]
    blob = packager("R-3 lab report.pdf", fake_pdf("3"))
    book = make_xlsx(tmp_path / "emb.xlsx", rows, [(3, 4, "oleObject1.bin", blob)])  # anchored on record 3
    sd = parse_document_bytes("doc", book.name, book.read_bytes(), cache_dir=tmp_path / "c")
    (child,) = sd.embedded
    (row,) = [r for r in table_rows(sd) if r.values[0] == "3"]
    assert (child.anchors[0].table_ref, child.anchors[0].table_row) == (row.table_ref, row.row_index)


def test_a_cached_parse_from_before_the_fix_is_glued_on_load(tmp_path, monkeypatch):
    from rag_wright.capabilities import sheet_tables

    book = _book(tmp_path / "db.xlsx", [H, ["1", "Jersey", 180, "pass"], [], ["2", "Rib", 220, "fail"]])
    real = sheet_tables.merge_continuations
    monkeypatch.setattr(sheet_tables, "merge_continuations", lambda document, data: False)  # the fix not there yet
    before = _rows(book, tmp_path / "c")
    assert [r.values[0] for r in before] == ["1"]  # record 2 read as a header and lost
    monkeypatch.setattr(sheet_tables, "merge_continuations", real)
    after = _rows(book, tmp_path / "c")  # the same cache, read after the fix
    assert [r.values[0] for r in after] == ["1", "2"] and len({r.table_ref for r in after}) == 1


def test_a_glued_database_passes_the_structural_eval_with_every_record_its_own_unit(tmp_path):
    book = _book(tmp_path / "db.xlsx", [H, ["1", "Jersey", 180, "pass"], ["2", "Rib", 220, "fail"], [],
                                        ["3", "Pique", 200, "pass"], [], ["4", "Fleece", 300, "pass"]])
    ev = evaluate_ingestion([IngestSource(path=str(book), table_mode="record")], cache_dir=tmp_path / "e")
    assert ev.passed, ev.failures
    assert ev.warnings == [] and ev.documents[0].tables == 1 and ev.documents[0].units == 4


def test_a_row_that_does_not_fill_the_columns_of_the_rows_above_is_not_one_of_them(tmp_path):
    block = [["Buyer", "Acme"], ["Style", "T-100"], ["Season", "SS26"], [], ["Construction notes", ""],
             ["Gauge", "28G"]]
    book = _book(tmp_path / "kv.xlsx", block)
    rows = _rows(book, tmp_path / "c")
    assert len({r.table_ref for r in rows}) == 2
    assert evaluate_ingestion([str(book)], cache_dir=tmp_path / "e").warnings == []
